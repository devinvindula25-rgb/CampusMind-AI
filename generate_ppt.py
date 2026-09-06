from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Helper to add a slide with a title and content
def add_slide(title, content_list):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    content_placeholder = slide.placeholders[1]
    
    title_placeholder.text = title
    tf = content_placeholder.text_frame
    tf.clear()
    
    for i, content in enumerate(content_list):
        p = tf.add_paragraph()
        p.text = content
        p.level = 0
        if i == 0:
            p.space_before = Pt(10)
        else:
            p.space_before = Pt(5)

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0] # Title Slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "CampusMind AI\nEmpowering Academic Well-being & Productivity"
subtitle.text = "Module: CMP 7003 - PRES1\nPresentation by [Your Name]\nStudent ID: [Your ID]"

# Slide 2: Introduction & Problem Definition
add_slide("Introduction & Problem Definition", [
    "The Academic Workload Challenge:",
    "• Academics juggle teaching, research, and high administrative burdens.",
    "• Existing tools are highly fragmented (emails, spreadsheets, separate calendars).",
    "• This leads to high cognitive load and severe risk of burnout.",
    "",
    "The Solution: CampusMind AI",
    "• A unified mobile dashboard tailored for faculty members.",
    "• Integrates scheduling, real-time burnout tracking, and AI-assisted task management."
])

# Slide 3: System Architecture
add_slide("System Architecture & Tech Stack", [
    "Frontend Framework:",
    "• React Native with Expo (cross-platform deployment).",
    "• Expo Router for modern, file-based routing and navigation.",
    "",
    "Backend & Real-time Database:",
    "• Firebase Authentication for secure user sessions.",
    "• Cloud Firestore NoSQL database for real-time synchronization.",
    "",
    "Intelligent Features:",
    "• Gemini AI API for generating contextual email drafts and well-being insights."
])

# Slide 4: Mobile SDK Integration (Learning Outcomes)
add_slide("Mobile SDK & Device Integration", [
    "Fulfilling key mobile platform learning outcomes:",
    "",
    "1. Calendar SDK (expo-calendar):",
    "   • Syncs university academic planners directly with the user's native device calendar.",
    "2. Persistent Storage (AsyncStorage):",
    "   • Caches user preferences and custom theme settings for offline access.",
    "3. Multimedia (expo-image):",
    "   • Optimized rendering for parallax graphics and responsive UI components.",
    "4. Haptics (expo-haptics):",
    "   • Tactile feedback during critical UI interactions."
])

# Slide 5: UI/UX & Design Principles
add_slide("UI/UX & Design Principles", [
    "Designed to reduce cognitive load:",
    "",
    "• Custom Theme Engine (SettingsContext): Fully supports Light and Dark modes.",
    "• Color Psychology: Activities and burnout risks are color-coded (e.g., Red = High Risk, Green = Low Risk).",
    "• Glassmorphism UI: Clean, subtle overlays that highlight critical data without overwhelming the user.",
    "• Unified Dashboard: All critical metrics (Schedule, Workload, Emails) accessible in one tap."
])

# Slide 6: Application Demonstration
add_slide("Demonstration of Application Features", [
    "Live Walkthrough of Key Features:",
    "",
    "1. Secure Authentication & Onboarding (Firebase Auth).",
    "2. Unified Dashboard: Real-time schedule sync and quick statistics.",
    "3. Academic Planner: Color-coded scheduling linked to device calendars.",
    "4. Burnout Monitor: Interactive charts (react-native-chart-kit) tracking stress levels.",
    "5. AI Email Assistant: Automated drafting of student responses using Gemini API.",
    "6. Floating AI Button: Contextual help accessible from any screen."
])

# Slide 7: Justification of Design Decisions
add_slide("Justification of Design Decisions", [
    "Why Expo Router over standard React Navigation?",
    "• File-based routing mathematically maps the directory structure, ensuring a highly scalable and maintainable codebase.",
    "",
    "Why Firebase Firestore?",
    "• Academics need instant updates across devices. Firestore's real-time NoSQL sync guarantees the dashboard is always current.",
    "",
    "Why React Context API?",
    "• Used for global Auth and Theme state. Avoided Redux to keep the architecture lightweight and highly performant."
])

# Slide 8: Critical Evaluation (Performance & Security)
add_slide("Critical Evaluation", [
    "Performance:",
    "• Used useMemo and useCallback React hooks to prevent heavy dashboard recalculations (e.g., burnout score).",
    "• App maintains a smooth 60fps even with real-time data streaming.",
    "",
    "Security:",
    "• Strict Firestore Security Rules ensure users can only access their specific UID data.",
    "• Firebase Auth tokens secure all backend communications.",
    "",
    "Scalability:",
    "• The serverless Firebase backend scales automatically with user traffic."
])

# Slide 9: Summary & Future Enhancements
add_slide("Summary & Future Enhancements", [
    "Summary:",
    "• Delivered a scalable, intelligent mobile application that directly addresses academic workload and well-being.",
    "",
    "Future Enhancements:",
    "1. Geolocation Integration (Maps):",
    "   • Navigational routing for new faculty members to find lecture halls on campus.",
    "2. Multimedia Upgrades:",
    "   • Integration of expo-av to allow attaching audio voice notes to research project logs.",
    "3. Deeper LMS Integration:",
    "   • Syncing directly with Moodle/Canvas."
])

# Slide 10: Conclusion
slide_layout = prs.slide_layouts[0] # Title Slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Thank You"
subtitle.text = "Any Questions?"

prs.save('CampusMind_AI_Presentation.pptx')
print("Presentation saved as CampusMind_AI_Presentation.pptx")
